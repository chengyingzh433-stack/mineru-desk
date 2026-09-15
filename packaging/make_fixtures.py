"""Generate public synthetic fixtures; no user documents are included."""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)
im = Image.new('RGB', (500, 220), 'white')
draw = ImageDraw.Draw(im)
draw.rectangle((30, 30, 190, 190), fill=(56, 118, 92))
draw.ellipse((290, 30, 450, 190), fill=(197, 133, 62))
im.save(out / 'figure.png')
d = Document()
d.add_heading('Physics Education', 0)
d.add_paragraph('Synthetic document for MinerU Desk. No personal data.')
t = d.add_table(rows=2, cols=2)
for cell, text in zip([t.cell(0,0),t.cell(0,1),t.cell(1,0),t.cell(1,1)], ['Method','Sample','Interview','12']):
    cell.text = text
d.add_picture(str(out / 'figure.png'), width=Inches(4))
d.add_paragraph('Figure 1. Synthetic shapes for image retention verification.')
d.save(out / 'selftest.docx')
c = canvas.Canvas(str(out / 'selftest.pdf'), pagesize=A4)
c.setTitle('Physics Education - MinerU Desk synthetic test')
for page in (1, 2):
    c.setFont('Helvetica-Bold', 24)
    c.drawString(50, 775, 'Physics Education')
    c.setFont('Helvetica', 12)
    c.drawString(50, 740, 'Synthetic PDF for local conversion testing. No personal data.')
    c.drawString(50, 700, 'Method: Interview. Sample: 12 participants.')
    c.drawImage(str(out / 'figure.png'), 50, 470, width=400, height=176)
    c.drawString(50, 440, 'Figure 1. Shapes used to verify image retention.')
    c.drawString(50, 55, f'Page {page} of 2')
    c.showPage()
c.save()
